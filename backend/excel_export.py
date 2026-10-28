"""
Builds the downloadable .xlsx reports:
  - build_report_workbook: one office's proposal, mirroring the province's
    existing format (office header, grouped columns, classification
    subtotals, grand total).
  - build_summary_workbook: a consolidated report totalling every office's
    proposal per account, for a given budget cycle.
"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

CLASSIFICATION_LABELS = {
    "PS": "PERSONAL SERVICES",
    "MOOE": "MAINTENANCE AND OTHER OPERATING EXPENDITURES",
    "FE": "FINANCIAL EXPENSES",
    "CO": "CAPITAL OUTLAY",
    "SPA": "SPECIAL PURPOSE APPROPRIATIONS",
    # Not an account classification: the 20% Development Fund project part of
    # a supplemental line (PPDO), reported as its own section.
    "DEV20": "20% DEVELOPMENT FUND PROJECTS",
}
CLASSIFICATION_ORDER = ["PS", "MOOE", "FE", "CO", "SPA", "DEV20"]


def split_dev20_lines(lines: list) -> list:
    """Splits supplemental lines that carry a 20% Development Fund project
    portion (`spa_portion`) into a regular line and a DEV20 line, so the
    projects are reported in their own section."""
    out = []
    for l in lines:
        p = l.get("spa_portion") or 0
        if p <= 0:
            out.append(l)
            continue
        prop = l.get("proposed_amount") or 0
        adj = l.get("adjusted_proposal")
        appr = l.get("approved_amount")
        dp = min(p, prop)
        regular = dict(l, proposed_amount=prop - dp,
                       adjusted_proposal=None if adj is None else adj - min(p, adj),
                       approved_amount=None if appr is None else appr - min(p, appr))
        regular["difference"] = regular["proposed_amount"] - (l.get("current_total") or 0)
        dev = dict(l, classification="DEV20", prev_year_actual=0, current_annual=0, current_supplementals=[],
                   current_total=0, proposed_amount=dp,
                   adjusted_proposal=None if adj is None else min(p, adj),
                   approved_amount=None if appr is None else min(p, appr),
                   difference=dp, remarks=None, remarks_adjusted=None, attachments=[])
        out.append(regular)
        out.append(dev)
    return out

def line_has_data(line: dict) -> bool:
    """True if a proposal line has any non-zero amount, a remark, or an
    attached document -- the reports leave out accounts with none of these."""
    amounts = [line.get("prev_year_actual"), line.get("current_annual"), line.get("current_total"),
               line.get("proposed_amount"), line.get("adjusted_proposal"), line.get("approved_amount")]
    amounts += [s.get("amount") for s in (line.get("current_supplementals") or [])]
    if any(a for a in amounts if a):
        return True
    if (line.get("remarks") or "").strip() or (line.get("remarks_adjusted") or "").strip():
        return True
    return bool(line.get("attachments"))


def summary_row_has_data(row: dict) -> bool:
    keys = ("total_prev_year_actual", "total_current_annual", "total_current_total",
            "total_proposed", "total_adjusted", "total_approved")
    return any(row.get(k) for k in keys)


THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
SECTION_FILL = PatternFill("solid", fgColor="EFEFEF")
CURRENCY_FMT = '#,##0.00;(#,##0.00);"-"'


def build_report_workbook(proposal: dict) -> io.BytesIO:
    """
    proposal: dict shaped like schemas.ProposalOut.
    """
    proposal = dict(proposal, lines=split_dev20_lines(proposal["lines"]))
    year = proposal["year"]
    prev_year = year - 2
    current_year = year - 1

    # How many "current year supplemental" reference columns does this
    # proposal actually use? Look at the widest line.
    max_supp = 0
    for line in proposal["lines"]:
        for s in line.get("current_supplementals", []):
            max_supp = max(max_supp, s["supplemental_number"])
    supp_col_count = max(max_supp, 1)  # always show at least one supplemental column

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget Report"

    col_prev = 3
    col_annual = 4
    col_supp_start = 5
    col_supp_end = col_supp_start + supp_col_count - 1
    col_total = col_supp_end + 1
    col_proposed = col_total + 1
    col_diff = col_proposed + 1
    col_docs = col_diff + 1
    col_remarks = col_docs + 1
    col_adjusted = col_remarks + 1
    col_remarks_adj = col_adjusted + 1
    col_approved = col_remarks_adj + 1
    ncols = col_approved

    # ---- Title row ----
    title_bits = [proposal["office_name"]]
    meta = []
    if proposal.get("office_code"):
        meta.append(proposal["office_code"])
    if proposal.get("office_sector"):
        meta.append(proposal["office_sector"])
    if meta:
        title_bits.append(f"({', '.join(meta)})")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=" ".join(title_bits))
    c.font = Font(name="Arial", bold=True, size=13)
    c.alignment = Alignment(horizontal="center")

    if proposal["budget_type"] == "annual":
        subtitle = f"ANNUAL BUDGET PROPOSAL - CY {year}"
    else:
        supp_no = proposal.get("supplemental_number") or 1
        subtitle = f"SUPPLEMENTAL BUDGET NO. {supp_no} PROPOSAL - CY {year}"
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=1, value=subtitle)
    c.font = Font(name="Arial", bold=True, size=11)
    c.alignment = Alignment(horizontal="center")

    # ---- Header rows ----
    header_row_top = 4
    header_row_bot = 5

    def single_header(col, top_text):
        ws.merge_cells(start_row=header_row_top, start_column=col, end_row=header_row_bot, end_column=col)
        ws.cell(row=header_row_top, column=col, value=top_text)

    single_header(1, "Object of Expenditures")
    single_header(2, "Account Code")
    single_header(col_prev, f"{prev_year} Actual")
    single_header(col_annual, f"Current Year {current_year}")
    for i in range(supp_col_count):
        single_header(col_supp_start + i, f"{current_year} Supplemental Budget No.{i + 1}")
    single_header(col_total, f"{current_year} Total")
    single_header(col_proposed, "Proposed")
    single_header(col_diff, "Difference")
    single_header(col_docs, "Supporting Documents")
    single_header(col_remarks, "Remarks for Proposal")
    single_header(col_adjusted, "Adjusted Proposal")
    single_header(col_remarks_adj, "Remarks for Adjusted Proposal")
    single_header(col_approved, "Approved")

    for row in (header_row_top, header_row_bot):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(name="Arial", bold=True, size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = HEADER_FILL
            cell.border = BORDER

    numeric_cols = [col_prev, col_annual] + list(range(col_supp_start, col_supp_end + 1)) + \
        [col_total, col_proposed, col_diff, col_adjusted, col_approved]
    text_cols = [col_docs, col_remarks, col_remarks_adj]

    # ---- Body ----
    row_idx = header_row_bot + 1
    grand_totals = {c: 0.0 for c in numeric_cols}

    lines_by_class = {k: [] for k in CLASSIFICATION_ORDER}
    for line in proposal["lines"]:
        if not line_has_data(line):
            continue  # leave out accounts with nothing against them
        lines_by_class.setdefault(line["classification"], []).append(line)

    for classification in CLASSIFICATION_ORDER:
        lines = lines_by_class.get(classification, [])
        if not lines:
            continue
        lines = sorted(lines, key=lambda l: (l["account_code"], l["account_name"]))

        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=ncols)
        cell = ws.cell(row=row_idx, column=1, value=CLASSIFICATION_LABELS[classification])
        cell.font = Font(name="Arial", bold=True, size=10)
        cell.fill = SECTION_FILL
        cell.border = BORDER
        row_idx += 1

        section_totals = {c: 0.0 for c in numeric_cols}
        for line in lines:
            supp_by_num = {s["supplemental_number"]: s["amount"] for s in line.get("current_supplementals", [])}
            approved = line.get("approved_amount")
            adjusted = line.get("adjusted_proposal")
            attachment_names = "; ".join(a["filename"] for a in line.get("attachments", []))

            row_values = {1: line["account_name"], 2: line["account_code"]}
            row_values[col_prev] = line["prev_year_actual"]
            row_values[col_annual] = line["current_annual"]
            for i in range(supp_col_count):
                row_values[col_supp_start + i] = supp_by_num.get(i + 1, 0)
            row_values[col_total] = line["current_total"]
            row_values[col_proposed] = line["proposed_amount"]
            row_values[col_diff] = line["difference"]
            row_values[col_docs] = attachment_names
            row_values[col_remarks] = line.get("remarks") or ""
            row_values[col_adjusted] = adjusted if adjusted is not None else ""
            row_values[col_remarks_adj] = line.get("remarks_adjusted") or ""
            row_values[col_approved] = approved if approved is not None else ""

            for col in range(1, ncols + 1):
                cell = ws.cell(row=row_idx, column=col, value=row_values.get(col, ""))
                cell.font = Font(name="Arial", size=9)
                cell.border = BORDER
                if col in numeric_cols:
                    cell.number_format = CURRENCY_FMT
                    cell.alignment = Alignment(horizontal="right", vertical="top")
                elif col in text_cols or col == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

            for col in numeric_cols:
                val = row_values.get(col)
                if isinstance(val, (int, float)):
                    section_totals[col] += val
            row_idx += 1

        subtotal_label = f"Total {CLASSIFICATION_LABELS[classification]}"
        ws.cell(row=row_idx, column=1, value=subtotal_label).font = Font(name="Arial", bold=True, size=9)
        ws.cell(row=row_idx, column=2).border = BORDER
        for col in range(3, ncols + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.border = BORDER
            cell.fill = SECTION_FILL
            if col in numeric_cols:
                cell.value = section_totals[col]
                cell.font = Font(name="Arial", bold=True, size=9)
                cell.number_format = CURRENCY_FMT
                cell.alignment = Alignment(horizontal="right")
                grand_totals[col] += section_totals[col]
        ws.cell(row=row_idx, column=1).border = BORDER
        ws.cell(row=row_idx, column=1).fill = SECTION_FILL
        row_idx += 1

    # ---- Grand total ----
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
    cell = ws.cell(row=row_idx, column=1, value="GRAND TOTAL")
    cell.font = Font(name="Arial", bold=True, size=10)
    cell.fill = HEADER_FILL
    cell.border = BORDER
    ws.cell(row=row_idx, column=2).border = BORDER
    ws.cell(row=row_idx, column=2).fill = HEADER_FILL
    for col in range(3, ncols + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.border = BORDER
        cell.fill = HEADER_FILL
        if col in numeric_cols:
            cell.value = grand_totals[col]
            cell.font = Font(name="Arial", bold=True, size=10)
            cell.number_format = CURRENCY_FMT
            cell.alignment = Alignment(horizontal="right")

    # ---- Column widths ----
    default_widths = {1: 42, 2: 14, col_prev: 14, col_annual: 16}
    for i in range(supp_col_count):
        default_widths[col_supp_start + i] = 18
    default_widths.update({
        col_total: 16, col_proposed: 14, col_diff: 16, col_docs: 28,
        col_remarks: 26, col_adjusted: 16, col_remarks_adj: 26, col_approved: 14,
    })
    for col in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(col)].width = default_widths.get(col, 14)
    ws.row_dimensions[header_row_top].height = 30
    ws.row_dimensions[header_row_bot].height = 42

    ws.freeze_panes = ws.cell(row=header_row_bot + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_summary_workbook(report: dict) -> io.BytesIO:
    """
    report: dict shaped like schemas.SummaryReportOut -- one row per account,
    totalled across every office for the given budget cycle.
    """
    year = report["year"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary Report"

    ncols = 9  # Object, Code, PrevActual, CurrAnnual, CurrTotal, Proposed, Adjusted, Approved, # Offices

    title = f"CONSOLIDATED BUDGET SUMMARY — ALL OFFICES — CY {year}"
    if report["budget_type"] == "supplemental":
        title = f"CONSOLIDATED SUPPLEMENTAL BUDGET NO. {report.get('supplemental_number') or 1} SUMMARY — ALL OFFICES — CY {year}"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(name="Arial", bold=True, size=13)
    c.alignment = Alignment(horizontal="center")

    subtitle = f"{report['offices_count']} office(s) included"
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=1, value=subtitle)
    c.font = Font(name="Arial", size=10)
    c.alignment = Alignment(horizontal="center")

    header_row = 4
    headers = [
        "Object of Expenditures", "Account Code", f"{year - 2} Actual (Total)",
        f"Current Year {year - 1} (Total)", f"{year - 1} Total (incl. Supplementals)",
        "Total Proposed (All Offices)", "Total Adjusted (All Offices)",
        "Total Approved (All Offices)", "# Offices with a Proposal",
    ]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=h)
        cell.font = Font(name="Arial", bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER

    row_idx = header_row + 1
    grand = [0.0] * 6
    rows_by_class = {k: [] for k in CLASSIFICATION_ORDER}
    for r in report["rows"]:
        if not summary_row_has_data(r):
            continue  # leave out accounts with nothing against them
        rows_by_class.setdefault(r["classification"], []).append(r)

    for classification in CLASSIFICATION_ORDER:
        rows = rows_by_class.get(classification, [])
        if not rows:
            continue
        rows = sorted(rows, key=lambda r: (r["account_code"], r["account_name"]))

        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=ncols)
        cell = ws.cell(row=row_idx, column=1, value=CLASSIFICATION_LABELS[classification])
        cell.font = Font(name="Arial", bold=True, size=10)
        cell.fill = SECTION_FILL
        cell.border = BORDER
        row_idx += 1

        section_totals = [0.0] * 6
        for r in rows:
            values = [
                r["account_name"], r["account_code"],
                r["total_prev_year_actual"], r["total_current_annual"], r["total_current_total"],
                r["total_proposed"], r["total_adjusted"], r["total_approved"], r["offices_included"],
            ]
            for col, val in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.font = Font(name="Arial", size=9)
                cell.border = BORDER
                if 3 <= col <= 8:
                    cell.number_format = CURRENCY_FMT
                    cell.alignment = Alignment(horizontal="right", vertical="top")
                elif col == 9:
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                elif col == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            for i, key in enumerate(["total_prev_year_actual", "total_current_annual", "total_current_total",
                                      "total_proposed", "total_adjusted", "total_approved"]):
                section_totals[i] += r[key]
            row_idx += 1

        ws.cell(row=row_idx, column=1, value=f"Total {CLASSIFICATION_LABELS[classification]}").font = Font(name="Arial", bold=True, size=9)
        ws.cell(row=row_idx, column=2).border = BORDER
        ws.cell(row=row_idx, column=1).border = BORDER
        ws.cell(row=row_idx, column=1).fill = SECTION_FILL
        for i, col in enumerate(range(3, 9)):
            cell = ws.cell(row=row_idx, column=col, value=section_totals[i])
            cell.font = Font(name="Arial", bold=True, size=9)
            cell.number_format = CURRENCY_FMT
            cell.alignment = Alignment(horizontal="right")
            cell.border = BORDER
            cell.fill = SECTION_FILL
            grand[i] += section_totals[i]
        ws.cell(row=row_idx, column=9).border = BORDER
        ws.cell(row=row_idx, column=9).fill = SECTION_FILL
        row_idx += 1

    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
    cell = ws.cell(row=row_idx, column=1, value="GRAND TOTAL")
    cell.font = Font(name="Arial", bold=True, size=10)
    cell.fill = HEADER_FILL
    cell.border = BORDER
    ws.cell(row=row_idx, column=2).border = BORDER
    ws.cell(row=row_idx, column=2).fill = HEADER_FILL
    for i, col in enumerate(range(3, 9)):
        cell = ws.cell(row=row_idx, column=col, value=grand[i])
        cell.font = Font(name="Arial", bold=True, size=10)
        cell.number_format = CURRENCY_FMT
        cell.alignment = Alignment(horizontal="right")
        cell.border = BORDER
        cell.fill = HEADER_FILL
    ws.cell(row=row_idx, column=9).border = BORDER
    ws.cell(row=row_idx, column=9).fill = HEADER_FILL

    widths = [42, 14, 16, 16, 16, 18, 18, 16, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[header_row].height = 36
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ----------------------------------------------------------- LBP Form No. 1 --

LOCAL_TAX_CAT = "Local Sources - Tax Revenue"
LOCAL_NONTAX_CAT = "Local Sources - Non-Tax Revenue"
BEGINNING_CASH_CAT = "Beginning Cash Balance"
NON_INCOME_CAT = "Non-Income Receipts"


def build_lbp_form1_workbook(data: dict) -> io.BytesIO:
    """
    Builds a consolidated 'Budget of Expenditures and Sources of Financing'
    report (LBP Form No. 1 layout), combining:
      - fund_sources: list of {category, particulars, amount}
      - summary_rows: per-account totals across all offices (same shape as
        schemas.SummaryAccountRow), used for the Expenditures section
      - spa_rows: list of {name, amount, auto_computed}
      - year, budget_type, supplemental_number, province_name

    Note: unlike the original government template, this doesn't split the
    current year into semesters (Actual/Estimate) -- this system doesn't
    collect budget data at that granularity, so the current-year column
    here is a single total.
    """
    year = data["year"]
    prev_year = year - 2
    current_year = year - 1
    province_name = data.get("province_name") or "___________________"

    fund_sources = data.get("fund_sources", [])
    summary_rows = data.get("summary_rows", [])
    spa_rows = data.get("spa_rows", [])

    wb = Workbook()
    ws = wb.active
    ws.title = "LBP Form 1"
    ncols = 5  # Particulars | Account Code | Prev Year Actual | Current Year Total | Budget Year Proposed

    def title_row(row, text, bold=True, size=11, center=True):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        c = ws.cell(row=row, column=1, value=text)
        c.font = Font(name="Arial", bold=bold, size=size)
        if center:
            c.alignment = Alignment(horizontal="center")
        return row + 1

    r = 1
    r = title_row(r, "LBP Form No. 1", size=10)
    r += 1
    r = title_row(r, "BUDGET OF EXPENDITURES AND SOURCES OF FINANCING", size=13)
    r = title_row(r, f"Province of {province_name}" if province_name else "PROVINCE", size=11)
    type_label = "ANNUAL" if data["budget_type"] == "annual" else f"SUPPLEMENTAL NO. {data.get('supplemental_number') or 1}"
    r = title_row(r, f"{type_label} BUDGET - CY {year}", size=10, bold=False)
    r += 1

    header_row = r
    headers = ["PARTICULARS", "Account Code", f"PAST YEAR {prev_year}\n(Actual)",
               f"CURRENT YEAR {current_year}\n(Total)", f"BUDGET YEAR {year}\n(PROPOSED)"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=h)
        cell.font = Font(name="Arial", bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
    ws.row_dimensions[header_row].height = 32
    r = header_row + 1

    def section_row(row, text):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        c = ws.cell(row=row, column=1, value=text)
        c.font = Font(name="Arial", bold=True, size=10)
        c.fill = SECTION_FILL
        c.border = BORDER
        return row + 1

    def data_row(row, name, code, prev, curr, proposed, bold=False, fill=None):
        vals = [name, code, prev, curr, proposed]
        for col, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = Font(name="Arial", bold=bold, size=9)
            cell.border = BORDER
            if col >= 3:
                cell.number_format = CURRENCY_FMT
                cell.alignment = Alignment(horizontal="right", vertical="top")
            elif col == 1:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill
        return row + 1

    # ---- I. Beginning Cash Balance ----
    r = section_row(r, "I. BEGINNING CASH BALANCE")
    beginning_total = 0.0
    for fs in fund_sources:
        if fs["category"] == BEGINNING_CASH_CAT:
            r = data_row(r, fs["particulars"], "", "", "", fs["amount"])
            beginning_total += fs["amount"]

    # ---- II. Receipts ----
    r = section_row(r, "II. RECEIPTS")
    r = section_row(r, "A. LOCAL SOURCES")
    local_total = 0.0
    for label, cat in [("1. Tax Revenue", LOCAL_TAX_CAT), ("2. Non-Tax Revenue", LOCAL_NONTAX_CAT)]:
        r = data_row(r, label, "", "", "", "", bold=True)
        sub = 0.0
        for fs in fund_sources:
            if fs["category"] == cat:
                r = data_row(r, "    " + fs["particulars"], "", "", "", fs["amount"])
                sub += fs["amount"]
        r = data_row(r, f"Total {label.split('. ')[1]}", "", "", "", sub, bold=True, fill=SECTION_FILL)
        local_total += sub
    r = data_row(r, "TOTAL LOCAL SOURCES", "", "", "", local_total, bold=True, fill=HEADER_FILL)

    r = section_row(r, "B. EXTERNAL SOURCES")
    external_total = 0.0
    for fs in fund_sources:
        if fs["category"].startswith("External Sources"):
            r = data_row(r, fs["particulars"], "", "", "", fs["amount"])
            external_total += fs["amount"]
    r = data_row(r, "TOTAL EXTERNAL SOURCES", "", "", "", external_total, bold=True, fill=HEADER_FILL)

    r = section_row(r, "C. NON-INCOME RECEIPTS")
    non_income_total = 0.0
    for fs in fund_sources:
        if fs["category"] == NON_INCOME_CAT:
            r = data_row(r, fs["particulars"], "", "", "", fs["amount"])
            non_income_total += fs["amount"]
    r = data_row(r, "TOTAL NON-INCOME RECEIPTS", "", "", "", non_income_total, bold=True, fill=HEADER_FILL)

    # Supplemental budgets (LBP Form 8) have sources under their own headings
    # (New Revenue Sources, Savings, Realignment, Reversion ...). They are
    # listed by group and counted as receipts.
    known = {BEGINNING_CASH_CAT, LOCAL_TAX_CAT, LOCAL_NONTAX_CAT, NON_INCOME_CAT}
    other_cats = []
    for fs in fund_sources:
        if fs["category"] not in known and not fs["category"].startswith("External Sources") \
                and fs["category"] not in other_cats:
            other_cats.append(fs["category"])
    other_total = 0.0
    if other_cats:
        r = section_row(r, "D. SUPPLEMENTAL SOURCES (LBP FORM 8)")
        for cat in other_cats:
            r = data_row(r, cat, "", "", "", "", bold=True)
            sub = 0.0
            for fs in fund_sources:
                if fs["category"] == cat:
                    r = data_row(r, "    " + fs["particulars"], "", "", "", fs["amount"])
                    sub += fs["amount"]
            r = data_row(r, f"Total {cat}", "", "", "", sub, bold=True, fill=SECTION_FILL)
            other_total += sub
        r = data_row(r, "TOTAL SUPPLEMENTAL SOURCES", "", "", "", other_total, bold=True, fill=HEADER_FILL)

    total_receipts = beginning_total + local_total + external_total + non_income_total + other_total
    r = data_row(r, "TOTAL RECEIPTS", "", "", "", total_receipts, bold=True, fill=HEADER_FILL)
    r += 1

    # ---- III. Expenditures (from office proposals, by classification) ----
    r = section_row(r, "III. EXPENDITURES")
    rows_by_class = {k: [] for k in CLASSIFICATION_ORDER}
    for row_data in summary_rows:
        if not summary_row_has_data(row_data):
            continue  # leave out accounts with nothing against them
        rows_by_class.setdefault(row_data["classification"], []).append(row_data)

    expenditure_class_totals = {}
    for classification in CLASSIFICATION_ORDER:
        if classification in ("SPA", "DEV20"):
            continue  # reported under "D. Special Purpose Appropriations" below
        rows = rows_by_class.get(classification, [])
        if not rows:
            expenditure_class_totals[classification] = 0.0
            continue
        rows = sorted(rows, key=lambda x: (x["account_code"], x["account_name"]))
        r = section_row(r, CLASSIFICATION_LABELS[classification])
        class_total = 0.0
        for row_data in rows:
            r = data_row(r, row_data["account_name"], row_data["account_code"],
                         row_data["total_prev_year_actual"], row_data["total_current_total"],
                         row_data["total_proposed"])
            class_total += row_data["total_proposed"]
        r = data_row(r, f"Total {CLASSIFICATION_LABELS[classification]}", "", "", "", class_total,
                     bold=True, fill=SECTION_FILL)
        expenditure_class_totals[classification] = class_total

    # ---- D. Special Purpose Appropriations ----
    r = section_row(r, "D. SPECIAL PURPOSE APPROPRIATIONS (SPA's)")
    spa_total = 0.0
    dev20_rows = sorted(rows_by_class.get("DEV20", []), key=lambda x: (x["account_code"], x["account_name"]))
    spa_acct_rows = sorted(rows_by_class.get("SPA", []), key=lambda x: (x["account_code"], x["account_name"]))
    is_supp = data.get("budget_type") == "supplemental"
    for spa in spa_rows:
        # In a supplemental the automatic 20% / 5% are not set aside separately:
        # they are spent through the actual project lines / LDRRMF account line
        # listed below, so the automatic figure is left out when those exist.
        if is_supp and spa.get("auto_computed") == "dev_fund_20pct" and dev20_rows:
            continue
        if is_supp and spa.get("auto_computed") == "ldrrmf_5pct" and spa_acct_rows:
            continue
        r = data_row(r, spa["name"], "", "", "", spa["amount"])
        spa_total += spa["amount"]
    if dev20_rows:
        r = data_row(r, "20% DEVELOPMENT FUND PROJECTS", "", "", "", "", bold=True)
        dev_total = 0.0
        for row_data in dev20_rows:
            r = data_row(r, "    " + row_data["account_name"], row_data["account_code"], "", "", row_data["total_proposed"])
            dev_total += row_data["total_proposed"]
        r = data_row(r, "Total 20% Development Fund Projects", "", "", "", dev_total, bold=True, fill=SECTION_FILL)
        spa_total += dev_total
    # office lines under SPA-classified accounts belong in this section too
    for row_data in spa_acct_rows:
        r = data_row(r, row_data["account_name"], row_data["account_code"], "", "", row_data["total_proposed"])
        spa_total += row_data["total_proposed"]
    r = data_row(r, "Total Special Purpose Appropriations", "", "", "", spa_total, bold=True, fill=SECTION_FILL)

    total_expenditures = sum(expenditure_class_totals.values()) + spa_total
    r = data_row(r, "TOTAL EXPENDITURES", "", "", "", total_expenditures, bold=True, fill=HEADER_FILL)
    r += 1

    # ---- IV. Ending Balance ----
    ending_balance = round(total_receipts - total_expenditures, 2)
    r = data_row(r, "IV. ENDING BALANCE", "", "", "", ending_balance, bold=True, fill=HEADER_FILL)

    widths = [46, 14, 16, 16, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
