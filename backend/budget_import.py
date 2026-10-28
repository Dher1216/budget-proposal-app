"""
Importing a provincial "Programmed Appropriation and Obligation by Object
of Expenditure" (LBP Form No. 2 style) workbook into the budget system.

The workbook this was built against (and tested exhaustively against) has
one worksheet per office, each shaped like:

    Row with "<OFFICE NAME> (<code>, <SECTOR>)"   <- office header
    Row with column headers (Object of Expenditures / Account Code / ...)
    "1.1 Personal Services:" / "1.2 Maintenance and Other Operating
      Expenditures" / "2.0 Capital Outlay" / financial-expenses section
      headers, each followed by its account rows
    "Total ... ►" subtotal rows, "TOTAL APPROPRIATION ►" grand total

Any worksheet that doesn't match this shape (summary tabs, charts, SPA
breakdown sheets, etc.) is skipped automatically -- detection is generic
(no hardcoded sheet names), so this keeps working on next year's workbook
even if sheet names/order change, as long as the overall layout doesn't.

Column layout (0-indexed, as read by openpyxl with values_only=True):
    0 = section header (e.g. "1.1 Personal Services:") or office header
    1 = account/object name (most rows)
    2 = account/object name, used instead of col 1 on a few indented rows;
        also holds the "Total ... ►" label for subtotal/total rows
    3 = account code
    4 = past year actual (e.g. 2025 Actual)
    5,6 = current year first/second semester (not used individually)
    7 = current year total (e.g. 2026 Total)
    8 = budget year proposed amount (e.g. 2027 Proposed) -- the authoritative
        final figure, verified column-by-column against each sheet's own
        "TOTAL APPROPRIATION" grand total (see tests)
"""
import re
from typing import Optional

import openpyxl

HEADER_RE = re.compile(r'^(.*?)\s*\(([^,()]+),\s*([^()]+)\)\s*$')
TOTAL_RE = re.compile(r'^\s*TOTAL\b.*[►>]\s*$', re.IGNORECASE)


def _classify(label: str) -> Optional[str]:
    u = label.upper()
    if "PERSONAL SERVICES" in u:
        return "PS"
    if "MAINTENANCE" in u and "OPERATING" in u:
        return "MOOE"
    if "CAPITAL OUTLAY" in u:
        return "CO"
    if "FINANCIAL EXPENSE" in u:
        return "FE"
    if "SPECIAL PURPOSE" in u:
        # "Special Purpose Appropriations" (e.g. "3.0 SPECIAL PURPOSE
        # APPROPRIATIONS") is its own classification, same as the file --
        # not folded into whichever section happened to come right before
        # it (previously Capital Outlay, since SPA usually follows it).
        return "SPA"
    return None


def _is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _num(v) -> float:
    return float(v) if isinstance(v, (int, float)) else 0.0


def _parse_office_sheet(ws):
    """Returns (office_name, office_code, office_sector, lines, control_total)
    or (None, None, None, [], None) if this sheet doesn't look like an
    office budget sheet at all."""
    office_name = office_code = office_sector = None
    cur_class = None
    lines = []
    control_total = None

    for row in ws.iter_rows(values_only=True):
        a = row[0] if len(row) > 0 else None
        b = row[1] if len(row) > 1 else None
        c = row[2] if len(row) > 2 else None
        d = row[3] if len(row) > 3 else None
        e = row[4] if len(row) > 4 else None
        h = row[7] if len(row) > 7 else None
        i = row[8] if len(row) > 8 else None

        if office_name is None and isinstance(a, str) and a.strip():
            m = HEADER_RE.match(a.strip())
            if m:
                office_name = m.group(1).strip()
                office_code = m.group(2).strip()
                office_sector = m.group(3).strip()
                continue

        if isinstance(a, str) and a.strip():
            cls = _classify(a)
            if cls:
                cur_class = cls
                continue

        name = None
        if not _is_blank(b):
            name = str(b).strip()
        elif not _is_blank(c):
            name = str(c).strip()

        if name is None or _is_blank(d):
            if name and TOTAL_RE.match(name) and "APPROPRIATION" in name.upper():
                control_total = _num(i)
            continue

        if TOTAL_RE.match(name):
            if "APPROPRIATION" in name.upper():
                control_total = _num(i)
            continue

        lines.append({
            "classification": cur_class,
            "code": str(d).strip(),
            "name": name,
            "prev_year_actual": _num(e),
            "current_annual": _num(h),
            "proposed_amount": _num(i),
        })

    return office_name, office_code, office_sector, lines, control_total


def parse_workbook(file_bytes) -> dict:
    """Parses the whole uploaded workbook. Returns a dict with an 'offices'
    list (one entry per detected office sheet) and a 'skipped_sheets' list
    (every other sheet, for transparency -- these are the summary/report
    tabs that aren't per-office budget data)."""
    wb = openpyxl.load_workbook(file_bytes, data_only=True)

    offices = []
    skipped_sheets = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        office_name, office_code, office_sector, lines, control_total = _parse_office_sheet(ws)

        if not office_name or not lines:
            skipped_sheets.append(sheet_name)
            continue

        total_proposed = round(sum(l["proposed_amount"] for l in lines), 2)
        control_match = (
            control_total is not None and abs(total_proposed - control_total) < 1.0
        )
        missing_classification = sum(1 for l in lines if l["classification"] is None)

        offices.append({
            "sheet": sheet_name,
            "office_name_in_file": office_name,
            "office_code": office_code,
            "office_sector": office_sector,
            "line_count": len(lines),
            "total_proposed": total_proposed,
            "control_total": control_total,
            "control_match": control_match,
            "missing_classification_count": missing_classification,
            "lines": lines,
        })

    return {"offices": offices, "skipped_sheets": skipped_sheets}


def normalize_office_name(name: str) -> str:
    """Loose normalization used to auto-suggest a matching existing office:
    uppercase, collapse whitespace, strip a leading 'OFFICE OF THE'/'OFFICE
    OF' so e.g. 'Office of the Provincial Administrator' and 'Provincial
    Administrator's Office' have a better chance of lining up."""
    u = re.sub(r"\s+", " ", name.strip().upper())
    u = re.sub(r"^OFFICE OF (THE )?", "", u)
    u = u.rstrip(".")
    return u


def suggest_office_match(office_name_in_file: str, existing_offices: list) -> Optional[int]:
    """existing_offices: list of (id, name) tuples. Returns the best-guess
    office id, or None if nothing looks close enough to auto-select --
    the caller (admin, in the preview UI) always gets the final say."""
    target = normalize_office_name(office_name_in_file)
    for oid, name in existing_offices:
        if normalize_office_name(name) == target:
            return oid
    # fall back to a loose containment match only when it's unambiguous
    candidates = [
        oid for oid, name in existing_offices
        if target in normalize_office_name(name) or normalize_office_name(name) in target
    ]
    if len(candidates) == 1:
        return candidates[0]
    return None


# =====================================================================
# Supplemental budget workbook (LBP Form 9 + LBP Form 8)
# =====================================================================
# A supplemental budget file has no per-office sheets. Instead:
#   * "LBP Form 9" lists every office as a heading row with its account
#     lines underneath, then (after a "PROGRAMS/PROJECTS TO BE FUNDED UNDER
#     THE 20% NTA DEVELOPMENT PROGRAM" heading) the 20% Development Fund
#     project breakdown, then a single 5% LDRRMF line, then (optionally)
#     the National Wealth utilisation section.
#   * "LBP Form 8" is the Statement of Funding Sources.
# Only the first "Amount" column is read; anything to its right is ignored.

LDRRMF_ACCOUNT_NAME = "5% Local Disaster Risk Reduction and Management Fund (LDRRMF)"
DEV_FUND_BALANCE_CATEGORY = "Reversion - 20% Development Fund"   # unexpended 20% DF carried into a supplemental
NON_RECEIPT_CATEGORIES = ("Beginning Cash Balance", "Savings", "Realignment", "Reversion", DEV_FUND_BALANCE_CATEGORY)
NTA_CATEGORY = "External Sources - IRA / National Tax Allocation"

_GENERIC_OBJECTS = {"MOOE", "PS", "CAPITAL OUTLAY", "CO", "FE", ""}


def _to_float(v) -> float:
    if isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace(",", "").strip())
        except ValueError:
            return 0.0
    return 0.0


def classification_from_code(code: str) -> str:
    """Best guess at an account's classification from its object code."""
    c = str(code).strip()
    if c.startswith("501"):
        return "PS"
    if c.startswith("503"):
        return "FE"
    if c.startswith("107") or c.startswith("106") or c.startswith("108"):
        return "CO"
    return "MOOE"


def _find_sheet(wb, needle: str):
    for name in wb.sheetnames:
        if needle in re.sub(r"\s+", " ", name.lower()):
            return wb[name]
    return None


def _find_header(ws, max_scan=40):
    """Locate the column header row of Form 9; returns dict of column
    indexes (0-based) or None."""
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), 1):
        cells = [str(c).strip().lower() if isinstance(c, str) else "" for c in row]
        if any("account code" in c for c in cells):
            cols = {"row": r_idx}
            for i, c in enumerate(cells):
                if "account code" in c and "code" not in cols:
                    cols["code"] = i
                elif c.startswith("object of expenditure") and "object" not in cols:
                    cols["object"] = i
                elif c.startswith("particulars") and "purpose" not in cols:
                    cols["purpose"] = i
                elif c == "amount" and "amount" not in cols:
                    cols["amount"] = i
            if "code" in cols and "amount" in cols:
                cols.setdefault("object", cols["code"] - 1)
                cols.setdefault("purpose", 1)
                return cols
    return None


def _parse_form9(ws):
    cols = _find_header(ws)
    if not cols:
        raise ValueError("Could not find the column headings (Account Code / Amount) on the LBP Form 9 sheet.")
    code_c, amt_c, obj_c, pur_c = cols["code"], cols["amount"], cols["object"], cols["purpose"]

    zone = "office"
    cur_office = None
    ldrrmf_amount = 0.0
    control_total = None
    buckets = {}      # office name -> {"lines": {code: line}, ...}
    order = []
    dev20 = {}        # code -> line   (project breakdown of the 20% Development Fund)

    def add_line(target: dict, code, amount, obj, purpose):
        ln = target.setdefault(code, {"code": code, "amount": 0.0, "object_hint": "", "particulars": []})
        ln["amount"] += amount
        if obj and not ln["object_hint"]:
            ln["object_hint"] = obj
        if purpose and len(ln["particulars"]) < 6:
            ln["particulars"].append(purpose)

    for row in ws.iter_rows(min_row=cols["row"] + 1, values_only=True):
        c0 = row[0].strip() if len(row) > 0 and isinstance(row[0], str) else ""
        code_raw = row[code_c] if len(row) > code_c else None
        code = str(int(code_raw)) if isinstance(code_raw, (int, float)) and not isinstance(code_raw, bool) \
            else (str(code_raw).strip() if code_raw is not None else "")
        has_code = code.isdigit()
        amount = _to_float(row[amt_c]) if len(row) > amt_c else 0.0
        up = c0.upper()

        if up.startswith("TOTAL APPROPRIATION"):
            control_total = amount
            break
        if up.startswith("UNAPPROPRIATED"):
            continue
        if c0 and not has_code:
            if "20%" in up and "NTA" in up and up.startswith("PROGRAMS"):
                zone = "dev20"; continue
            if up.startswith("5%") and "DISASTER" in up:
                ldrrmf_amount += amount
                zone = "after5"; continue
            if up.startswith("PROGRAMS") and "NATIONAL WEALTH" in up:
                zone = "office"; cur_office = None; continue
            if zone == "office":
                cur_office = c0
                if c0 not in buckets:
                    buckets[c0] = {}
                    order.append(c0)
            continue   # sector / municipality / project-group headings in the 20% zone
        if has_code and amount:
            obj = row[obj_c].strip() if len(row) > obj_c and isinstance(row[obj_c], str) else ""
            purpose = row[pur_c].strip() if len(row) > pur_c and isinstance(row[pur_c], str) else ""
            if zone == "dev20":
                add_line(dev20, code, amount, obj, purpose)
            elif zone == "office" and cur_office:
                add_line(buckets[cur_office], code, amount, obj, purpose)
    return buckets, order, dev20, ldrrmf_amount, control_total


def _parse_form8(ws):
    """Grouped funding sources: main headings (1.0 New Revenue Sources, 3.0
    Savings, 4.0 Realignment, 5.0 Reversion ...) become the category. Only
    the top-level item lines are taken; the indented realignment/project
    detail underneath them is a breakdown of those items, not extra money."""
    heading_re = re.compile(r"^\d+(\.\d+)?\s+(.+?)\s*$")
    start = None
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=40, values_only=True), 1):
        if isinstance(row[0], str) and row[0].strip().lower().startswith("particulars"):
            start = r_idx
            break
    if start is None:
        raise ValueError("Could not find the 'Particulars' heading on the LBP Form 8 sheet.")
    amt_c = 4
    for row in ws.iter_rows(min_row=start, max_row=start, values_only=True):
        for i, c in enumerate(row):
            if isinstance(c, str) and c.strip().lower() == "amount":
                amt_c = i

    group = None
    items = []
    total_income = None
    for row in ws.iter_rows(min_row=start + 1, values_only=True):
        c0 = row[0].strip() if isinstance(row[0], str) else ""
        if c0.upper().startswith("TOTAL ESTIMATED INCOME"):
            total_income = _to_float(row[amt_c]); break
        m = heading_re.match(c0) if c0 else None
        if m and not c0.startswith("("):
            title = m.group(2).strip().title()
            group = title
            continue
        label = row[1].strip() if len(row) > 1 and isinstance(row[1], str) else ""
        amount = _to_float(row[amt_c]) if len(row) > amt_c else 0.0
        if group and label and label.lower() != "xxx" and amount:
            category = group
            if re.search(r"20\s*%\s*development\s*fund", label, re.I):
                category = DEV_FUND_BALANCE_CATEGORY
            elif group.lower().startswith("new revenue") and re.search(r"\bNTA\b|national tax|\bIRA\b", label, re.I):
                category = NTA_CATEGORY
            items.append({"category": category, "particulars": label, "amount": round(amount, 2)})
    return items, total_income


def parse_supplemental_workbook(file_bytes) -> dict:
    wb = openpyxl.load_workbook(file_bytes, data_only=True)
    ws9 = _find_sheet(wb, "form 9")
    if ws9 is None:
        raise ValueError("No 'LBP Form 9' sheet was found in this workbook.")
    buckets, order, dev20, ldrrmf_amount, control_total = _parse_form9(ws9)

    ws8 = _find_sheet(wb, "form 8")
    fund_sources, form8_total = ([], None)
    if ws8 is not None:
        fund_sources, form8_total = _parse_form8(ws8)

    # 20% Development Fund project breakdown belongs to PPDO; the 5% LDRRMF
    # line to PDRRMO. They're merged into that office's own entry when the
    # file already lists it; otherwise an entry is created for it.
    def find_or_add(token, default_name):
        for name in order:
            if token in name.upper():
                return name
        buckets[default_name] = {}
        order.append(default_name)
        return default_name

    ppdo_name = find_or_add("PLANNING", "OFFICE OF THE PROVINCIAL PLANNING & DEVELOPMENT COORDINATOR") if dev20 else None
    pdrrmo_name = find_or_add("DISASTER", "OFFICE OF THE PROVINCIAL DISASTER RISK REDUCTION AND MANAGEMENT OFFICE") if ldrrmf_amount else None

    dev20_total = round(sum(l["amount"] for l in dev20.values()), 2)
    if ppdo_name:
        for code, ln in dev20.items():
            tgt = buckets[ppdo_name].setdefault(code, {"code": code, "amount": 0.0, "object_hint": ln["object_hint"], "particulars": []})
            tgt["amount"] += ln["amount"]
            tgt["particulars"] = (tgt["particulars"] + ln["particulars"])[:6]
    if pdrrmo_name:
        buckets[pdrrmo_name]["LDRRMF"] = {
            "code": "LDRRMF", "amount": ldrrmf_amount, "object_hint": LDRRMF_ACCOUNT_NAME, "particulars": ["5% LDRRMF"],
        }

    offices = []
    grand = 0.0
    for name in order:
        lines = [dict(l, amount=round(l["amount"], 2)) for l in buckets[name].values() if l["amount"]]
        if not lines:
            continue
        total = round(sum(l["amount"] for l in lines), 2)
        grand += total
        offices.append({
            "office_name_in_file": name,
            "office_code": None,
            "office_sector": None,
            "line_count": len(lines),
            "total_amount": total,
            "dev20_included": dev20_total if name == ppdo_name else 0,
            "ldrrmf_included": round(ldrrmf_amount, 2) if name == pdrrmo_name else 0,
            "lines": lines,
        })
    return {
        "offices": offices,
        "grand_total": round(grand, 2),
        "control_total": control_total,
        "control_match": control_total is not None and abs(grand - control_total) < 1.0,
        "dev20_total": dev20_total,
        "ldrrmf_total": round(ldrrmf_amount, 2),
        "fund_sources": fund_sources,
        "fund_total": round(sum(f["amount"] for f in fund_sources), 2),
        "form8_total": form8_total,
        "form8_match": form8_total is not None and abs(sum(f["amount"] for f in fund_sources) - form8_total) < 1.0,
    }
